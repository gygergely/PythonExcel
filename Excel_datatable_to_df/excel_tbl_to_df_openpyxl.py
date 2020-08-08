import openpyxl as pyxl
import pandas as pd


def load_data_table_to_df(xl_file_name, sheet_name, table_name):
    """
    Load an Excel data table content to a pandas dataframe
    :param xl_file_name: Excel file path & name
    :param sheet_name: sheet name where the data table(s) can be found
    :param table_name: table name
    :return: pandas dataframe or in case of error None
    """
    try:
        # Assign the Excel file to a variable
        wb = pyxl.load_workbook(xl_file_name)
        # Get sheet
        ws = wb[sheet_name]
        # Get table
        data_tbl = ws.tables[table_name]
        # Get table range
        cells = ws[data_tbl.ref]

        # Iterate through table range cells and create a list of lists
        cells_value = list()

        for cell in cells:
            row = list()
            for i in range(0, len(cell)):
                row.append(cell[i].value)

            cells_value.append(row)

        # Load the list to a dataframe
        df = pd.DataFrame(cells_value)
        # Grab the first row for the header
        df_header = df.iloc[0]
        # Get the data except the 1st row
        df = df[1:]
        # Set the 1st row as header
        df.columns = df_header
        # Reset df index
        df.reset_index(drop=True, inplace=True)

        return df

    except Exception as ex:
        template = "An exception of type {0} occurred. Arguments:\n{1!r}"
        message = template.format(type(ex).__name__, ex.args)
        print(message)


if __name__ == '__main__':

    df_test1 = load_data_table_to_df('data_tbl_test.xlsx', 'test_datatbl_1', 'tbl_test_3')
    df_test2 = load_data_table_to_df('data_tbl_test.xlsx', 'test_datatbl_2', 'tbl_test_1')
    df_test3 = load_data_table_to_df('data_tbl_test.xlsx', 'test_datatbl_2', 'tbl_test_2')

    if df_test1 is not None:
        print(df_test1.head())

    if df_test2 is not None:
        print(df_test2.head())

    if df_test3 is not None:
        print(df_test3.head())
